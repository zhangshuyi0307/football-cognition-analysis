#!/usr/bin/env python3
"""Rebuild non-MOT reports with one canonical HR3 roster.

Canonical samples (not physiology-driven behavioral exclusions):
* Corsi/SART/Flanker/Soccer: 69 = athlete 24 + non-athlete high HR 23 +
  non-athlete low HR 22.
* BELT: 68, because one high-HR non-athlete has no raw BELT record.

The script preserves all calculable results, including non-significant tests.
"""
from __future__ import annotations

import itertools
import json
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy import stats
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

F = Path(r"F:\足球实验数据\分析输出")
OUT = F / "各范式重建报告_指定N69_BELT68_20260729"
TABLES = OUT / "表格"
FIGS = OUT / "图表"
MASTER = F / "other_five_behavior_tasks_updated_quality_cleaning" / "other_five_tasks_participant_metrics_masked.csv"
ROSTER = F / "行为分析被试名册_手机号桥接补入" / "行为分析被试名册.csv"
HR = F / "PMwG_遗漏运动员1825补入审计_20260729" / "PMwG_66人_心率三组输入_含补入运动员.csv"
SART = F / "SART黄色注意捕获_两组分析" / "SART_被试条件级指标.csv"
SOCCER = F / "足球决策题_运动员答案众数" / "足球决策正式trial_脱敏.csv"
SURVEY = F / "survey_annotated_final_three_groups" / "survey_responses_annotated_athlete_groups.csv"

GROUPS = ["运动员", "非运动员高心率", "非运动员低心率"]
COLORS = {"运动员": "#3F6FB5", "非运动员高心率": "#E18727", "非运动员低心率": "#368C63"}
TASKS = {
    "Corsi": "Corsi反向空间工作记忆",
    "Flanker": "Flanker选择性注意",
    "BELT": "BELT风险学习",
}
LABELS = {
    "accuracy_rate": "正确率",
    "max_correct_sequence_length": "最大正确序列长度",
    "mean_first_click_rt_ms": "平均首次点击RT(ms)",
    "mean_response_duration_ms": "平均反应时长(ms)",
    "cash_rate": "兑现率",
    "explosion_rate": "爆炸率",
    "mean_pump_count": "平均充气次数",
    "mean_trial_duration_ms": "平均trial时长(ms)",
    "total_score_final": "总得分",
    "flanker_interference_rt_ms": "Flanker冲突RT代价(ms)",
    "mean_rt_correct_congruent_ms": "一致条件正确RT(ms)",
    "mean_rt_correct_incongruent_ms": "不一致条件正确RT(ms)",
    "mean_rt_correct_neutral_ms": "中性条件正确RT(ms)",
    "mean_rt_correct_ms": "总体正确RT(ms)",
    "dprime": "SART d′",
    "go_rt_mean_ms": "正确Go RT(ms)",
    "go_rt_cv": "正确Go RT CoV",
    "commission_rate": "NoGo误按率",
    "go_omission_rate": "Go漏按率",
}


def key(x: object) -> str:
    s = "" if pd.isna(x) else str(x)
    digits = "".join(c for c in s if c.isdigit())
    if len(digits) >= 7:
        return f"{digits[:3]}...{digits[-4:]}"
    return s.strip()


def ptext(p: float) -> str:
    return "NA" if not np.isfinite(p) else ("<.001" if p < .001 else f"{p:.3f}")


def mark(p: float) -> str:
    if not np.isfinite(p):
        return "无法估计"
    return "显著" if p < .05 else ("趋势" if p < .10 else "不显著")


def sem(x: pd.Series) -> float:
    x = pd.to_numeric(x, errors="coerce").dropna()
    return float(x.std(ddof=1) / math.sqrt(len(x))) if len(x) > 1 else math.nan


def cliffs(a: np.ndarray, b: np.ndarray) -> float:
    return float((np.sum(a[:, None] > b) - np.sum(a[:, None] < b)) / (len(a) * len(b)))


def setup() -> None:
    for path in [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"]:
        if Path(path).exists():
            from matplotlib import font_manager
            font_manager.fontManager.addfont(path)
            plt.rcParams["font.family"] = font_manager.FontProperties(fname=path).get_name()
            break
    plt.rcParams["axes.unicode_minus"] = False
    TABLES.mkdir(parents=True, exist_ok=True)
    FIGS.mkdir(parents=True, exist_ok=True)


def canonical_groups() -> pd.DataFrame:
    roster = pd.read_csv(ROSTER, encoding="utf-8-sig")
    roster["k"] = roster["masked_phone"].map(key)
    athletes = roster.loc[roster["group_two"].eq("运动员"), ["k"]].drop_duplicates().copy()
    athletes["hr3"] = "运动员"
    h = pd.read_csv(HR, encoding="utf-8-sig")
    h["k"] = h["masked_phone"].map(key)
    non = h.loc[h["hr_group3"].isin(["非运动员高心率", "非运动员低心率"]), ["k", "hr_group3"]].drop_duplicates("k")
    non = non.rename(columns={"hr_group3": "hr3"})
    out = pd.concat([athletes, non], ignore_index=True).drop_duplicates("k", keep="first")
    out = out[out["hr3"].isin(GROUPS)].copy()
    assert out.groupby("hr3").size().to_dict() == {"运动员": 24, "非运动员高心率": 23, "非运动员低心率": 22}, out.groupby("hr3").size().to_dict()
    out.to_csv(TABLES / "心率三组_69人规范名册_脱敏.csv", index=False, encoding="utf-8-sig")
    return out


def tests(frame: pd.DataFrame, value: str, condition: str = "") -> tuple[dict, list[dict]]:
    vals = {g: pd.to_numeric(frame.loc[frame.hr3.eq(g), value], errors="coerce").dropna().to_numpy(float) for g in GROUPS}
    usable = [vals[g] for g in GROUPS if len(vals[g]) >= 2]
    if len(usable) != 3:
        h, p, overall_reason = math.nan, math.nan, "insufficient_group_data"
    elif np.ptp(np.concatenate(usable)) == 0:
        h, p, overall_reason = 0.0, 1.0, "all_identical_values_p_set_to_1"
    else:
        h, p, overall_reason = (*stats.kruskal(*usable), "")
    overall = {"条件": condition, "指标": LABELS.get(value, value), "n运动员": len(vals[GROUPS[0]]), "n高心率": len(vals[GROUPS[1]]), "n低心率": len(vals[GROUPS[2]]), "Kruskal_H": h, "原始p": p, "p原因": overall_reason, "标注": mark(p)}
    pairs = []
    for a, b in itertools.combinations(GROUPS, 2):
        x, y = vals[a], vals[b]
        if len(x) >= 1 and len(y) >= 1:
            if np.ptp(np.concatenate([x, y])) == 0:
                u, pp, d, pair_reason = len(x) * len(y) / 2, 1.0, 0.0, "all_identical_values_p_set_to_1"
            else:
                u, pp = stats.mannwhitneyu(x, y, alternative="two-sided")
                d, pair_reason = cliffs(x, y), ""
        else:
            u = pp = d = math.nan
            pair_reason = "insufficient_group_data"
        pairs.append({"条件": condition, "指标": LABELS.get(value, value), "组别A": a, "组别B": b, "nA": len(x), "nB": len(y), "均值A": np.mean(x) if len(x) else math.nan, "均值B": np.mean(y) if len(y) else math.nan, "MW_U": u, "原始p": pp, "p原因": pair_reason, "Cliffs_delta": d, "标注": mark(pp)})
    return overall, pairs


def add_doc_table(doc: Document, df: pd.DataFrame, title: str, limit: int = 30) -> None:
    doc.add_heading(title, level=2)
    if df.empty:
        doc.add_paragraph("无可估计结果。")
        return
    shown = df.head(limit).copy().fillna("")
    if len(df) > limit:
        doc.add_paragraph(f"正文显示前{limit}行；完整结果见同目录CSV。")
    table = doc.add_table(rows=1, cols=len(shown.columns))
    table.style = "Table Grid"
    for i, c in enumerate(shown.columns):
        table.rows[0].cells[i].text = str(c)
    for _, row in shown.iterrows():
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = f"{v:.3f}" if isinstance(v, float) else str(v)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7.5)
                    r.font.name = "Microsoft YaHei"
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def new_doc(title: str, sample_text: str) -> Document:
    d = Document()
    for style in ["Normal", "Heading 1", "Heading 2"]:
        d.styles[style].font.name = "Microsoft YaHei"
        d.styles[style]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(19)
    d.add_paragraph(sample_text)
    d.add_paragraph("统计口径：三组总体比较为Kruskal-Wallis；两两比较为Mann-Whitney U；报告原始p值及Cliff's delta。所有可计算结果均保留，不因不显著而删除。")
    return d


def add_image(d: Document, path: Path, caption: str) -> None:
    d.add_picture(str(path), width=Cm(15.7))
    p = d.add_paragraph(caption); p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def simple_plot(long: pd.DataFrame, metrics: list[str], title: str, prefix: str) -> Path:
    rows = math.ceil(len(metrics) / 2)
    fig, axes = plt.subplots(rows, 2, figsize=(14, max(5, 4.2 * rows)))
    axes = np.ravel(axes)
    for ax, m in zip(axes, metrics):
        q = long[long.metric.eq(m)]
        overall, _ = tests(q.rename(columns={"value": m}), m)
        for i, g in enumerate(GROUPS):
            v = pd.to_numeric(q.loc[q.hr3.eq(g), "value"], errors="coerce").dropna()
            x = np.full(len(v), i, dtype=float) + np.random.default_rng(20260729 + i).normal(0, .045, len(v))
            ax.scatter(x, v, s=23, alpha=.65, color=COLORS[g])
            if len(v):
                ax.errorbar(i, v.mean(), yerr=sem(v), fmt="o", color=COLORS[g], capsize=4, linewidth=2)
        ax.set_xticks(range(3), GROUPS, rotation=12)
        ax.set_title(f"{LABELS.get(m,m)} | Kruskal p={ptext(overall['原始p'])}")
        ax.grid(axis="y", alpha=.25)
    for ax in axes[len(metrics):]: ax.axis("off")
    handles = [plt.Line2D([0],[0], marker="o", color="w", markerfacecolor=COLORS[g], label=g, markersize=7) for g in GROUPS]
    fig.legend(handles=handles, loc="upper center", ncol=3, frameon=False)
    fig.suptitle(title, y=.995, fontsize=16)
    fig.tight_layout(rect=(0,0,1,.95))
    path = FIGS / f"{prefix}_核心指标_散点误差柱_原始p.png"
    fig.savefig(path, dpi=190, bbox_inches="tight", facecolor="white"); plt.close(fig)
    return path


def one_metric_plot(long: pd.DataFrame, metric: str, title: str, prefix: str) -> Path:
    """One metric per figure: raw points + mean/SEM + raw Kruskal p."""
    q = long[long.metric.eq(metric)].copy()
    q = q.rename(columns={"value": metric})
    overall, _ = tests(q, metric)
    fig, ax = plt.subplots(figsize=(7.8, 5.4))
    for i, group in enumerate(GROUPS):
        values = pd.to_numeric(q.loc[q.hr3.eq(group), metric], errors="coerce").dropna()
        x = np.full(len(values), i, dtype=float) + np.random.default_rng(20260730 + i).normal(0, .045, len(values))
        ax.scatter(x, values, color=COLORS[group], s=34, alpha=.68, zorder=2)
        if len(values):
            ax.errorbar(i, values.mean(), yerr=sem(values), fmt="o", color=COLORS[group], capsize=5, linewidth=2.2, markersize=8, zorder=3)
    label = LABELS.get(metric, metric)
    ax.set_xticks(range(3), GROUPS, rotation=10)
    ax.set_ylabel(label)
    ax.set_title(f"{title}: {label}\nKruskal-Wallis 原始 p={ptext(overall['原始p'])}")
    ax.grid(axis="y", alpha=.25)
    fig.tight_layout()
    path = FIGS / f"{prefix}_{metric}_散点误差柱_原始p.png"
    fig.savefig(path, dpi=190, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def build_metric_report(master: pd.DataFrame, paradigm: str) -> None:
    q = master[master.paradigm.eq(paradigm)].copy()
    q["metric"] = q.metric.astype(str)
    metrics = list(q.metric.drop_duplicates())
    desc = q.groupby(["hr3", "metric"], as_index=False).agg(n=("value","count"), mean=("value","mean"), sd=("value","std"), median=("value","median"), min=("value","min"), max=("value","max"))
    overall, pairs = [], []
    for m in metrics:
        z = q[q.metric.eq(m)].copy().rename(columns={"value": m})
        a, b = tests(z, m)
        overall.append(a); pairs.extend(b)
    overall, pairs = pd.DataFrame(overall), pd.DataFrame(pairs)
    tag = TASKS[paradigm]
    desc.to_csv(TABLES / f"{tag}_描述统计_69或68人.csv", index=False, encoding="utf-8-sig")
    overall.to_csv(TABLES / f"{tag}_Kruskal总体检验_完整.csv", index=False, encoding="utf-8-sig")
    pairs.to_csv(TABLES / f"{tag}_两两比较_完整.csv", index=False, encoding="utf-8-sig")
    # Keep a compact overview, but produce one standalone p-labelled chart per metric.
    overview = simple_plot(q, metrics, f"{tag}：心率三组核心指标", tag)
    figures = [one_metric_plot(q, metric, tag, tag) for metric in metrics]
    expected = 68 if paradigm == "BELT" else 69
    n_now = q.groupby("hr3")["k"].nunique().to_dict()
    d = new_doc(f"{tag}报告（重建版）", f"预设样本口径：{expected}人。实际指标覆盖：{n_now}。任务缺失不会被生理数据替代或人为填补。")
    d.add_heading("1. 样本与清洗", level=1)
    d.add_paragraph("本报告仅按任务自身的可用记录计算指标；心率三组仅用于分组比较，不作为删除行为数据的依据。BELT中爆炸trial是行为结果，不作为错误trial删除。")
    add_doc_table(d, desc, "表1. 描述统计", 60)
    d.add_heading("2. 三组比较", level=1)
    add_doc_table(d, overall, "表2. Kruskal-Wallis总体检验", 30)
    add_image(d, overview, "图1. 全部核心指标概览；散点、均值±SEM与三组原始p值。")
    for index, image in enumerate(figures, start=2):
        add_image(d, image, f"图{index}. {LABELS.get(metrics[index-2], metrics[index-2])}：全部个体散点、均值±SEM与原始p值。")
    add_doc_table(d, pairs, "表3. 两两Mann-Whitney比较", 60)
    d.add_paragraph("解释时应同时看总体检验、两两原始p值和Cliff's delta；由于多指标检验较多，本版将结果作为探索性描述。")
    d.save(OUT / f"{tag}_重建报告_20260729.docx")


def build_belt_questionnaire(master: pd.DataFrame) -> None:
    belt = master[master.paradigm.eq("BELT")].pivot_table(index="k", columns="metric", values="value", aggfunc="first").reset_index()
    survey = pd.read_csv(SURVEY, encoding="utf-8-sig")
    survey["k"] = survey["masked_phone_behavior"].fillna(survey["masked_phone"]).map(key)
    scales = [c for c in ["big5_extraversion","big5_agreeableness","big5_conscientiousness","big5_emotional_stability","big5_openness","grit_total","resilience_total","dassy_total","life_satisfaction_mean","loneliness_mean"] if c in survey]
    d = belt.merge(survey[["k", *scales]], on="k", how="left")
    rows = []
    bcols = [c for c in belt.columns if c != "k"]
    for b, s in itertools.product(bcols, scales):
        z = d[[b,s]].dropna()
        rho, p = stats.spearmanr(z[b], z[s]) if len(z) >= 4 else (math.nan, math.nan)
        rows.append({"BELT指标": LABELS.get(b,b), "问卷指标": s, "n":len(z), "Spearman_rho":rho, "原始p":p, "标注":mark(p)})
    out = pd.DataFrame(rows).sort_values("原始p")
    out.to_csv(TABLES / "BELT_问卷_Spearman相关_68人完整.csv", index=False, encoding="utf-8-sig")
    matrix = out.pivot(index="BELT指标", columns="问卷指标", values="Spearman_rho")
    fig, ax = plt.subplots(figsize=(13,5.5))
    im=ax.imshow(matrix, vmin=-.6, vmax=.6, cmap="coolwarm")
    ax.set_xticks(range(len(matrix.columns)), matrix.columns, rotation=42, ha="right"); ax.set_yticks(range(len(matrix.index)), matrix.index)
    for i in range(matrix.shape[0]):
        for j in range(matrix.shape[1]):
            v=matrix.iloc[i,j]
            p=out[(out["BELT指标"].eq(matrix.index[i]))&(out["问卷指标"].eq(matrix.columns[j]))]["原始p"].iloc[0]
            ax.text(j,i,f"{v:.2f}\np={ptext(p)}",ha="center",va="center",fontsize=7)
    fig.colorbar(im, ax=ax, label="Spearman rho"); ax.set_title("BELT与问卷相关热图（68人，原始p）"); fig.tight_layout()
    path=FIGS / "BELT_问卷相关热图_68人_原始p.png"; fig.savefig(path,dpi=190,bbox_inches="tight");plt.close(fig)
    doc=new_doc("BELT与问卷相关补充报告", "样本：有BELT行为记录且能桥接到问卷的被试。相关不等于因果；心率三组不作为BELT主问题的唯一解释。")
    add_doc_table(doc,out,"表4. BELT与问卷Spearman相关（全部）",60); add_image(doc,path,"图2. 相关热图；每格为rho及原始p。")
    doc.save(OUT / "BELT与问卷相关补充报告_68人_20260729.docx")


def build_sart(groups: pd.DataFrame) -> None:
    d = pd.read_csv(SART, encoding="utf-8-sig")
    d["k"] = d["masked_phone"].map(key)
    d = d.merge(groups, on="k", how="inner")
    metrics = [c for c in ["dprime","go_rt_mean_ms","go_rt_cv","commission_rate","go_omission_rate"] if c in d]
    desc = d.melt(id_vars=["k","hr3","phase","capture"], value_vars=metrics, var_name="metric", value_name="value")
    overall=[]; pairs=[]
    for (phase,capture,m), q in desc.groupby(["phase","capture","metric"]):
        z=q.rename(columns={"value":m});a,b=tests(z,m,f"{phase} / {capture}");overall.append(a);pairs.extend(b)
    overall,pairs=pd.DataFrame(overall),pd.DataFrame(pairs)
    desc.to_csv(TABLES / "SART_69人_被试条件指标_long.csv",index=False,encoding="utf-8-sig")
    overall.to_csv(TABLES / "SART_69人_Kruskal总体检验_完整.csv",index=False,encoding="utf-8-sig")
    pairs.to_csv(TABLES / "SART_69人_两两比较_完整.csv",index=False,encoding="utf-8-sig")
    sart_figures=[]
    for m in metrics:
        fig,axes=plt.subplots(1,2,figsize=(14,5.2))
        for ax,capture in zip(axes,["普通","黄色捕获"]):
            q=desc[(desc.metric.eq(m)) & (desc.capture.eq(capture))]
            for gi,g in enumerate(GROUPS):
                z=q[q.hr3.eq(g)]
                for i,phase in enumerate(["前测","后测"]):
                    raw=z[z.phase.eq(phase)].value.dropna()
                    ax.scatter(np.random.default_rng(20260730+gi*10+i).normal(i+(gi-1)*.18,.025,len(raw)),raw,color=COLORS[g],alpha=.42,s=17)
                    if len(raw): ax.errorbar(i+(gi-1)*.18,raw.mean(),yerr=sem(raw),fmt="o",color=COLORS[g],capsize=4,label=g if i==0 else None)
            pvals=overall[(overall["指标"].eq(LABELS.get(m,m))) & (overall["条件"].astype(str).str.contains(str(capture)))]["原始p"].tolist()
            ax.set_xticks([0,1],["前测","后测"]);ax.set_title(f"{capture} | 原始p=" + "/".join(ptext(x) for x in pvals));ax.grid(axis="y",alpha=.25)
        axes[0].set_ylabel(LABELS.get(m,m));axes[0].legend(frameon=False);fig.suptitle(f"SART {LABELS.get(m,m)}：前后测×捕获条件×心率三组",y=1.02);fig.tight_layout()
        path=FIGS/f"SART_69人_{m}_前后测捕获条件_散点误差柱_原始p.png";fig.savefig(path,dpi=190,bbox_inches="tight");plt.close(fig);sart_figures.append(path)
    doc=new_doc("SART黄色注意捕获报告（69人重建版）","69人：运动员24、非运动员高心率23、非运动员低心率22。SART错误保留为commission/omission结果；RT只来自正确Go。")
    add_doc_table(doc,overall,"表1. 条件级三组Kruskal-Wallis检验（完整）",60)
    for i,path in enumerate(sart_figures, start=1):
        add_image(doc,path,f"图{i}. {LABELS.get(metrics[i-1],metrics[i-1])}：各条件个体散点、均值±SEM与原始p值。")
    add_doc_table(doc,pairs,"表2. 条件级两两Mann-Whitney比较（完整）",60)
    doc.save(OUT / "SART黄色注意捕获报告_69人_20260729.docx")


def build_soccer(groups: pd.DataFrame) -> None:
    d=pd.read_csv(SOCCER,encoding="utf-8-sig")
    d["k"]=d["masked_phone"].map(key)
    # This athlete registered with 189****1104 but completed the behavioral
    # tasks under 182****6863.  The same verified alias is used by Flanker.
    soccer_groups = pd.concat(
        [groups, pd.DataFrame({"k": ["182...6863"], "hr3": ["运动员"]})],
        ignore_index=True,
    )
    d=d.merge(soccer_groups,on="k",how="inner")
    d=d[d.ans.astype(str).isin(list("ABCD")) & d.block.isin([1,2])].copy()
    coverage=d.groupby("hr3").k.nunique().reindex(GROUPS,fill_value=0)
    # Answer vector analysis has no assumed correct answer.
    rows=[]
    for block in [1,2]:
        q=d[d.block.eq(block)]
        for a,b in itertools.combinations(GROUPS,2):
            ca=q[q.hr3.eq(a)].ans.value_counts().reindex(list("ABCD"),fill_value=0).to_numpy();cb=q[q.hr3.eq(b)].ans.value_counts().reindex(list("ABCD"),fill_value=0).to_numpy()
            chi,p,_,_=stats.chi2_contingency(np.vstack([ca,cb]))
            n=ca.sum()+cb.sum();v=math.sqrt(chi/n) if n else math.nan
            rows.append({"Block":block,"组别A":a,"组别B":b,"nA":int(q[q.hr3.eq(a)].k.nunique()),"nB":int(q[q.hr3.eq(b)].k.nunique()),"chi2":chi,"原始p":p,"Cramers_V":v,"标注":mark(p)})
    tests_df=pd.DataFrame(rows);tests_df.to_csv(TABLES/"足球决策_69人_Block选择分布两两比较_完整.csv",index=False,encoding="utf-8-sig")
    props=d.groupby(["block","hr3","ans"]).size().rename("n").reset_index();props["ratio"]=props.groupby(["block","hr3"]).n.transform(lambda x:x/x.sum());props.to_csv(TABLES/"足球决策_69人_Block选择比例.csv",index=False,encoding="utf-8-sig")
    soccer_figures=[]
    for block in [1,2]:
        fig,ax=plt.subplots(figsize=(7.5,5.4))
        q=props[props.block.eq(block)];bottom=np.zeros(3)
        for ans,color in zip("ABCD",["#4C78A8","#F58518","#54A24B","#E45756"]):
            z=q[q.ans.eq(ans)].set_index("hr3").ratio.reindex(GROUPS,fill_value=0).to_numpy();ax.bar(GROUPS,z,bottom=bottom,label=ans,color=color);bottom+=z
        p0=tests_df[tests_df.Block.eq(block)]["原始p"].min();ax.set_title(f"Block {block} | 最小两两原始 p={ptext(p0)}");ax.set_xticks(range(3),GROUPS,rotation=12);ax.set_ylim(0,1);ax.set_ylabel("选择比例");ax.legend(title="选项",frameon=False);fig.tight_layout()
        path=FIGS/f"足球决策_69人_Block{block}_选择分布_原始p.png";fig.savefig(path,dpi=190,bbox_inches="tight");plt.close(fig);soccer_figures.append(path)
    fig,ax=plt.subplots(figsize=(9,5.2));plot=tests_df.copy();plot["比较"]=plot["组别A"]+" vs "+plot["组别B"]+" / B"+plot["Block"].astype(str)
    bars=ax.barh(plot["比较"],plot["Cramers_V"],color=["#C94C4C" if p<.05 else "#7F8C8D" for p in plot["原始p"]]);
    for bar,p in zip(bars,plot["原始p"]): ax.text(bar.get_width()+.005,bar.get_y()+bar.get_height()/2,f"p={ptext(p)}",va="center",fontsize=9)
    ax.set_xlabel("Cramer's V");ax.set_title("足球决策：每个Block的两两选择向量差异");ax.grid(axis="x",alpha=.25);fig.tight_layout();pairfig=FIGS/"足球决策_69人_Block两两向量差异_CramersV_原始p.png";fig.savefig(pairfig,dpi=190,bbox_inches="tight");plt.close(fig)
    doc=new_doc("足球决策选择向量报告（69人重建版）",f"按心率三组桥接后的原始选择分布覆盖：{coverage.to_dict()}。不使用原正确答案或运动员众数作为主结论，只比较A/B/C/D选择向量。")
    add_doc_table(doc,tests_df,"表1. Block内三组两两选择分布比较",30)
    for i,path in enumerate(soccer_figures,start=1): add_image(doc,path,f"图{i}. Block{i} 的A/B/C/D选择比例；标题为最小两两原始p。")
    add_image(doc,pairfig,"图3. 每个Block与各两两组别的Cramer's V；标签为原始p。")
    doc.add_paragraph("注：若原始足球trial无法完整桥接到69人，本报告会在覆盖表中如实显示，不将缺失行为记录补造为选择。")
    doc.save(OUT / "足球决策选择向量报告_69人_20260729.docx")
    coverage.rename("n").reset_index().to_csv(TABLES/"足球决策_69人_原始trial覆盖审计.csv",index=False,encoding="utf-8-sig")


def main() -> None:
    setup(); groups=canonical_groups()
    master=pd.read_csv(MASTER,encoding="utf-8-sig");master["k"]=master.masked_phone.map(key);master=master.merge(groups,on="k",how="inner")
    audit=[]
    for task in ["Corsi","BELT","Flanker"]:
        q=master[master.paradigm.eq(task)]
        audit.append({"范式":task,"预设人数":68 if task=="BELT" else 69,"实际唯一人数":q.k.nunique(),"分组人数":json.dumps(q.groupby("hr3").k.nunique().to_dict(),ensure_ascii=False)})
        build_metric_report(master,task)
    build_belt_questionnaire(master);build_sart(groups);build_soccer(groups)
    audit=pd.DataFrame(audit);audit.to_csv(TABLES/"各范式重建样本量审计.csv",index=False,encoding="utf-8-sig")
    (OUT/"README.md").write_text("# 重建报告\n\n本目录统一使用规范心率三组名册。Corsi/SART/Flanker/Soccer预设69人；BELT预设68人。每份报告旁均保留完整CSV与p值标注图。\n",encoding="utf-8")
    print(json.dumps({"output":str(OUT),"groups":groups.hr3.value_counts().to_dict()},ensure_ascii=False))


if __name__ == "__main__":
    main()
