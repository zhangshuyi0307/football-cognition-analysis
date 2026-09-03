from __future__ import annotations

import json
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
import matplotlib.pyplot as plt
from matplotlib import font_manager

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except Exception as exc:
    raise SystemExit(f"python-docx is required: {exc}")


F = Path(r"F:\足球实验数据")
SRC = F / "分析输出" / "MOT_DRT_resource_lite_HR3_20260730"
OUT = F / "分析输出" / "MOT_DRT_simple_proxy_HR3_66_stage_load_20260730"
FIG = OUT / "figures"
TAB = OUT / "tables"
STAGES = ["rest", "exercise1", "exercise2"]
STAGE_LABEL = {"rest": "静息", "exercise1": "第一次运动后", "exercise2": "第二次运动后"}
GROUPS = ["athlete", "nonathlete_high_hr", "nonathlete_low_hr"]
GROUP_LABEL = {
    "athlete": "运动员",
    "nonathlete_high_hr": "非运动员高心率",
    "nonathlete_low_hr": "非运动员低心率",
}
METRICS = [
    ("mot_accuracy", "MOT正确率"),
    ("mot_dprime", "MOT d′"),
    ("drt_hit_rate", "DRT命中率"),
    ("drt_miss_rate", "DRT漏报率"),
    ("drt_logrt", "DRT logRT"),
    ("G_lite", "整体资源代理 G_lite"),
    ("pi_MOT_lite", "MOT优先资源代理 pi_MOT_lite"),
]


def pick_font():
    for name in ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Arial"]:
        matches = [f for f in font_manager.findSystemFonts() if name.lower() in Path(f).stem.lower()]
        if matches:
            return name
    return "DejaVu Sans"


def p_label(p):
    if pd.isna(p):
        return "NA"
    if p < 0.001:
        return "p<.001"
    return f"p={p:.3f}"


def cliffs_delta(a, b):
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    a = a[np.isfinite(a)]
    b = b[np.isfinite(b)]
    if not len(a) or not len(b):
        return np.nan
    return float((np.greater.outer(a, b).sum() - np.less.outer(a, b).sum()) / (len(a) * len(b)))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)
    TAB.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = pick_font()
    plt.rcParams["axes.unicode_minus"] = False

    src = SRC / "resource_lite_subject_stage_load.csv"
    df = pd.read_csv(src)
    df["drt_miss_rate"] = 1 - pd.to_numeric(df["drt_hit_rate"], errors="coerce")
    df["mot_accuracy"] = pd.to_numeric(df["selected"], errors="coerce") / pd.to_numeric(df["target_total"], errors="coerce")
    for col, _ in METRICS:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = df[df["target_count"].isin([1, 2, 3, 4, 5]) & df["stage"].isin(STAGES) & df["group"].isin(GROUPS)].copy()

    # This report deliberately contains no pooled stage/load group test.
    desc_rows, kw_rows, pair_rows = [], [], []
    for stage in STAGES:
        for load in [1, 2, 3, 4, 5]:
            cell = df[(df.stage == stage) & (df.target_count == load)]
            for key, label in METRICS:
                vals = {g: cell.loc[cell.group == g, key].dropna().to_numpy() for g in GROUPS}
                for g in GROUPS:
                    x = vals[g]
                    desc_rows.append({
                        "stage": stage, "阶段": STAGE_LABEL[stage], "target_count": load,
                        "球数": load, "metric": key, "指标": label, "group": g,
                        "组别": GROUP_LABEL[g], "n": len(x),
                        "mean": np.mean(x) if len(x) else np.nan,
                        "sd": np.std(x, ddof=1) if len(x) > 1 else np.nan,
                        "median": np.median(x) if len(x) else np.nan,
                        "q1": np.quantile(x, .25) if len(x) else np.nan,
                        "q3": np.quantile(x, .75) if len(x) else np.nan,
                        "min": np.min(x) if len(x) else np.nan,
                        "max": np.max(x) if len(x) else np.nan,
                    })
                nonempty = [vals[g] for g in GROUPS if len(vals[g])]
                try:
                    h, p = stats.kruskal(*nonempty) if len(nonempty) == 3 and all(len(v) >= 2 for v in nonempty) else (np.nan, np.nan)
                except Exception:
                    h, p = np.nan, np.nan
                kw_rows.append({
                    "stage": stage, "阶段": STAGE_LABEL[stage], "target_count": load, "球数": load,
                    "metric": key, "指标": label, "n_total": sum(len(v) for v in nonempty),
                    "n_athlete": len(vals["athlete"]), "n_high_hr": len(vals["nonathlete_high_hr"]),
                    "n_low_hr": len(vals["nonathlete_low_hr"]), "Kruskal_H": h, "raw_p": p,
                    "显著性": "显著" if pd.notna(p) and p < .05 else ("趋势" if pd.notna(p) and p < .10 else ("不显著" if pd.notna(p) else "NA")),
                })
                for i, ga in enumerate(GROUPS):
                    for gb in GROUPS[i + 1:]:
                        a, b = vals[ga], vals[gb]
                        try:
                            u, pp = stats.mannwhitneyu(a, b, alternative="two-sided") if len(a) >= 2 and len(b) >= 2 else (np.nan, np.nan)
                        except Exception:
                            u, pp = np.nan, np.nan
                        pair_rows.append({
                            "stage": stage, "阶段": STAGE_LABEL[stage], "target_count": load, "球数": load,
                            "metric": key, "指标": label, "group_a": ga, "组别A": GROUP_LABEL[ga],
                            "group_b": gb, "组别B": GROUP_LABEL[gb], "n_a": len(a), "n_b": len(b),
                            "U": u, "raw_p": pp, "Cliffs_delta_A_minus_B": cliffs_delta(a, b),
                            "显著性": "显著" if pd.notna(pp) and pp < .05 else ("趋势" if pd.notna(pp) and pp < .10 else ("不显著" if pd.notna(pp) else "NA")),
                        })

    desc = pd.DataFrame(desc_rows)
    kw = pd.DataFrame(kw_rows)
    pairs = pd.DataFrame(pair_rows)
    desc.to_csv(TAB / "MOT_DRT_描述统计_阶段×球数×三组_完整.csv", index=False, encoding="utf-8-sig")
    kw.to_csv(TAB / "MOT_DRT_Kruskal_每阶段×每球数_三组_完整.csv", index=False, encoding="utf-8-sig")
    pairs.to_csv(TAB / "MOT_DRT_MannWhitney_每阶段×每球数_三组两两_完整.csv", index=False, encoding="utf-8-sig")
    shutil.copy2(src, TAB / "resource_lite_subject_stage_load_原始代理数据.csv")

    # One figure per stage x load x metric; no pooled condition figure is used as the primary result.
    fig_rows = []
    colors = {"athlete": "#2f6db3", "nonathlete_high_hr": "#e07a27", "nonathlete_low_hr": "#31915e"}
    for stage in STAGES:
        for load in [1, 2, 3, 4, 5]:
            cell = df[(df.stage == stage) & (df.target_count == load)]
            for key, label in METRICS:
                kwrow = kw[(kw.stage == stage) & (kw.target_count == load) & (kw.metric == key)].iloc[0]
                fig, ax = plt.subplots(figsize=(8.6, 5.8), dpi=180)
                for idx, g in enumerate(GROUPS):
                    x = cell.loc[cell.group == g, key].dropna().to_numpy()
                    xpos = idx + 1
                    if len(x):
                        jitter = np.linspace(-.12, .12, len(x)) if len(x) > 1 else np.array([0])
                        ax.scatter(np.full(len(x), xpos) + jitter, x, s=20, alpha=.42, color=colors[g], edgecolor="none")
                        mean = np.mean(x)
                        se = np.std(x, ddof=1) / np.sqrt(len(x)) if len(x) > 1 else 0
                        ax.errorbar(xpos, mean, yerr=se, fmt="o", color=colors[g], ecolor=colors[g], capsize=4, lw=2, ms=7, zorder=5)
                ax.set_xticks([1, 2, 3], [GROUP_LABEL[g] for g in GROUPS])
                ax.set_title(f"{label}：{STAGE_LABEL[stage]}，{load}球")
                ax.set_ylabel(label)
                ax.grid(axis="y", alpha=.25)
                ax.text(.02, .98, f"三组 Kruskal-Wallis {p_label(kwrow.raw_p)}", transform=ax.transAxes, va="top", fontsize=10)
                ax.set_xlim(.5, 3.5)
                fig.tight_layout()
                fn = f"{stage}_load{load}_{key}_三组组间_raw_p.png"
                fig.savefig(FIG / fn, bbox_inches="tight")
                plt.close(fig)
                fig_rows.append({"stage": stage, "阶段": STAGE_LABEL[stage], "target_count": load, "球数": load, "metric": key, "指标": label, "图文件": fn, "Kruskal_raw_p": kwrow.raw_p})
    pd.DataFrame(fig_rows).to_csv(TAB / "图表索引_每阶段×每球数×指标.csv", index=False, encoding="utf-8-sig")

    # Report: tables are complete; figures are inserted in compact metric sections.
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"; normal.font.size = Pt(9)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MOT/DRT 三组简易资源代理池报告\n阶段×球数完整组间比较")
    run.bold = True; run.font.size = Pt(18); run.font.name = "Microsoft YaHei"
    doc.add_paragraph("本版本严格按每个阶段×每个球数分别比较三组，不进行跨阶段或跨球数的总体组间比较。资源指标为行为代理，不是 PMwG 后验潜变量。")
    doc.add_heading("一、样本与指标", level=1)
    doc.add_paragraph("共同纳入样本 N=66：运动员 21 人、非运动员高心率 23 人、非运动员低心率 22 人。每个条件都使用该条件实际可用的被试数，并在表格中逐项报告。")
    doc.add_paragraph("简易代理定义：G_lite=(标准化 MOT d′ + 标准化 DRT 综合效率)/2；pi_MOT_lite=标准化 MOT d′ − 标准化 DRT 综合效率。DRT 综合效率由命中率和负 logRT 构成。")
    doc.add_heading("二、统计口径", level=1)
    doc.add_paragraph("每个阶段×球数条件进行三组 Kruskal-Wallis 总体检验；随后进行三组两两 Mann-Whitney U 检验，并报告 Cliff's delta。全部为原始 p 值，不做 FDR。显著：p<.05；趋势：.05≤p<.10。")
    doc.add_paragraph("完整 CSV 表格位于 tables 文件夹；每一张条件图均标注对应的三组 Kruskal-Wallis 原始 p 值，两两 p 值见两两比较表。")
    doc.add_heading("三、完整结果", level=1)
    for key, label in METRICS:
        doc.add_heading(label, level=2)
        sub = kw[kw.metric == key].copy()
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for c, text in zip(table.rows[0].cells, ["阶段", "球数", "n", "Kruskal H", "原始 p"]): c.text = text
        for _, row in sub.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row["阶段"]); cells[1].text = str(int(row["球数"]))
            cells[2].text = f"{int(row.n_athlete)}/{int(row.n_high_hr)}/{int(row.n_low_hr)}"
            cells[3].text = f"{row.Kruskal_H:.3f}" if pd.notna(row.Kruskal_H) else "NA"
            cells[4].text = p_label(row.raw_p)
        doc.add_paragraph("图表：下面按阶段和球数分别给出散点+均值±SEM图。")
        for stage in STAGES:
            for load in [1, 2, 3, 4, 5]:
                fn = next(x["图文件"] for x in fig_rows if x["stage"] == stage and x["target_count"] == load and x["metric"] == key)
                p = doc.add_paragraph(f"{STAGE_LABEL[stage]}，{load}球")
                p.runs[0].bold = True
                doc.add_picture(str(FIG / fn), width=Inches(5.8))
    doc.add_heading("四、解释边界", level=1)
    doc.add_paragraph("G_lite 和 pi_MOT_lite 是由已清洗的 MOT/DRT 行为指标构成的简易代理，不能称为原文 PMwG 的潜在 G 或 pi。正式 PMwG 继续单独运行；本报告用于先行查看每个阶段×球数的三组行为差异。")
    doc.save(OUT / "MOT_DRT_简易资源代理池_阶段×球数三组完整报告_20260730.docx")

    manifest = {
        "n": 66, "groups": {"athlete": 21, "nonathlete_high_hr": 23, "nonathlete_low_hr": 22},
        "stages": STAGES, "loads": [1, 2, 3, 4, 5], "metrics": [x[0] for x in METRICS],
        "primary_test": "Kruskal-Wallis within each stage x load; no pooled stage/load group test",
        "pairwise_test": "Mann-Whitney U within each stage x load",
        "figures": len(fig_rows), "tables": [p.name for p in TAB.glob("*.csv")],
    }
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "README.md").write_text("# MOT/DRT 简易资源代理池\n\n本目录严格按阶段×球数分别进行三组比较，不进行跨阶段或跨球数总体组间比较。\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
